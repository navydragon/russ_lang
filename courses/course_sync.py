"""Добавление курса и заданий без очистки таблиц и потери результатов."""
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
import re

from django.core.management.base import CommandError
from openpyxl import load_workbook

from courses.models import Course, Lesson, Task
from courses.task_catalog import TASK_CATALOG


COURSE_NAME = 'Русский язык для иностранных студентов'
DEFAULT_TASKS_FILE = Path(__file__).parent / 'Книга русский.xlsx'


@dataclass(frozen=True)
class DesiredTask:
    lesson: int
    code: str
    name: str
    legacy_code: str = ''


@dataclass
class SyncPlan:
    course_id: int | None
    lessons: dict
    missing_lessons: list
    new_tasks: list = field(default_factory=list)
    code_updates: list = field(default_factory=list)
    existing_tasks: int = 0
    regular_count: int = 0


def read_regular_tasks(path):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise CommandError(f'Не удалось прочитать файл заданий: {path}') from exc
    tasks = []
    try:
        for line, row in enumerate(workbook.active.iter_rows(min_row=2, values_only=True), 2):
            if not row or not row[0]:
                continue
            code = str(row[0]).strip()
            match = re.match(r'^task(\d+)-', code)
            if not match:
                raise CommandError(f'Строка {line}: не распознан номер урока в {code}.')
            lesson = int(match.group(1))
            # Старые id уроков из Excel не используем как первичные ключи БД.
            if not 1 <= lesson <= 10 or len(row) < 3 or row[2] != lesson:
                raise CommandError(f'Строка {line}: противоречивый номер урока для {code}.')
            name = str(row[1]).strip() if row[1] else f'Задание {code.rsplit("-", 1)[-1]}'
            if len(code) > 100 or len(name) > 200:
                raise CommandError(f'Строка {line}: превышена длина кода или названия.')
            tasks.append(DesiredTask(lesson, code, name))
    finally:
        workbook.close()
    if not tasks:
        raise CommandError('Файл не содержит заданий.')
    duplicates = [code for code, count in Counter(t.code for t in tasks).items() if count > 1]
    if duplicates:
        raise CommandError(f'Повторяющиеся коды в Excel: {duplicates}')
    return tasks


def build_plan(regular_tasks):
    courses = list(Course.objects.filter(name=COURSE_NAME))
    if len(courses) > 1:
        raise CommandError('Найдено несколько курсов с целевым названием; нужна проверка дублей.')
    course = courses[0] if courses else None
    lessons = {l.position: l for l in Lesson.objects.filter(course=course)} if course else {}
    missing_lessons = [n for n in range(1, 11) if n not in lessons]
    plan = SyncPlan(course.pk if course else None, lessons, missing_lessons, regular_count=len(regular_tasks))
    ai_tasks = [DesiredTask(t.lesson, t.code, t.name, t.legacy_code) for t in TASK_CATALOG]
    desired = [*regular_tasks, *ai_tasks]
    codes = {t.code for t in desired} | {t.legacy_code for t in ai_tasks}
    found = defaultdict(list)
    for task in Task.objects.filter(code__in=codes).select_related('lesson'):
        found[task.code].append(task)
    positions = defaultdict(int)
    if course:
        for task in Task.objects.filter(lesson__course=course).select_related('lesson'):
            positions[task.lesson.position] = max(positions[task.lesson.position], task.position)
    for definition in desired:
        matches = found[definition.code]
        legacy = [t for t in found[definition.legacy_code] if t.lesson.course_id == plan.course_id] if definition.legacy_code else []
        if len(matches) > 1 or len(legacy) > 1 or (matches and legacy):
            raise CommandError(f'Неоднозначные записи для кода {definition.code}.')
        existing = matches[0] if matches else (legacy[0] if legacy else None)
        if existing:
            if existing.lesson.course_id != plan.course_id or existing.lesson.position != definition.lesson:
                raise CommandError(f'Код {definition.code} занят заданием другого курса или урока.')
            if legacy:
                plan.code_updates.append((existing.pk, existing.code, definition.code))
            plan.existing_tasks += 1
            continue
        # Новые задания добавляются после существующих; их позиции не меняются.
        positions[definition.lesson] += 1
        plan.new_tasks.append((definition, positions[definition.lesson]))
    return plan


def apply_plan(plan):
    course = Course.objects.get(pk=plan.course_id) if plan.course_id else Course.objects.create(name=COURSE_NAME)
    lessons = dict(plan.lessons)
    for number in plan.missing_lessons:
        lessons[number] = Lesson.objects.create(course=course, position=number, name=f'Урок {number}')
    for pk, old_code, new_code in plan.code_updates:
        changed = Task.objects.filter(pk=pk, code=old_code).update(code=new_code)
        if changed != 1:
            raise CommandError(f'Запись {old_code} изменилась во время синхронизации.')
    for definition, position in plan.new_tasks:
        Task.objects.create(
            lesson=lessons[definition.lesson], position=position,
            code=definition.code, name=definition.name,
        )
    return course
