from io import BytesIO
from unittest.mock import patch, Mock
from urllib.error import HTTPError, URLError

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import OperationalError, connection
from django.db.migrations.executor import MigrationExecutor
from django.test import TestCase, TransactionTestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from courses.models import Course, Lesson, Task, TaskAttempt
from . import portal, roster
from .forms import GroupForm
from .importing import HEADERS, ImportError, read_rows, template_bytes
from .models import Group, Student


def workbook(rows):
    book = Workbook()
    book.active.append(HEADERS)
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def row(code='001', **changes):
    data = dict(last_name='Иванов', first_name='Иван', middle_name='', email='test@example.test', code=code, password='private-password')
    data.update(changes)
    return data


def remote(action, data, method='POST'):
    if action == 'upsert_group':
        return {'group_id': '800', 'status': 'unchanged'}
    if action == 'upsert_student':
        return {'student_id': '900', 'status': 'updated'}
    if action == 'ensure_membership':
        return {'student_id': '900', 'group_id': '800', 'status': 'unchanged'}
    return {'status': 'removed'}


class RosterTests(TestCase):
    def setUp(self):
        self.group = Group.objects.create(name='Русский', code='RU01')
        self.staff = get_user_model().objects.create_user(username='staff', is_staff=True)
        self.client.force_login(self.staff)

    def student(self):
        return Student.objects.create(code='001', last_name='Старый', first_name='Иван', group=self.group)

    def file(self, rows=None):
        data = rows or [['Иванов', 'Иван', '', 'test@example.test', '001', '001', 'private-password']]
        return SimpleUploadedFile('students.xlsx', workbook(data))

    @patch('students.portal.call', side_effect=remote)
    def test_upsert_preserves_local_identity_and_results(self, api):
        student = self.student()
        course = Course.objects.create(name='Курс')
        lesson = Lesson.objects.create(name='Урок', course=course)
        task = Task.objects.create(name='Задание', lesson=lesson)
        attempt = TaskAttempt.objects.create(student=student, task=task, datetime=timezone.now(), results='saved')
        for _ in range(2):
            obj, status = roster.save_student(self.group.pk, row())
            self.assertEqual(obj.pk, student.pk)
        roster.remove_student(self.group.pk, student.pk)
        roster.remove_student(self.group.pk, student.pk)
        student.refresh_from_db()
        self.assertIsNone(student.group_id)
        self.assertEqual(student.email, 'test@example.test')
        self.assertEqual(Student.objects.count(), 1)
        self.assertTrue(TaskAttempt.objects.filter(pk=attempt.pk, results='saved').exists())

    @patch('students.portal.call', side_effect=portal.PortalError('timeout', uncertain=True))
    def test_timeout_does_not_commit(self, api):
        with self.assertRaises(portal.PortalError):
            roster.save_student(self.group.pk, row())
        self.assertFalse(Student.objects.exists())

    @patch('students.portal.call', side_effect=remote)
    def test_lost_response_then_safe_retry(self, api):
        calls = 0
        def uncertain(action, data, method='POST'):
            nonlocal calls
            calls += 1
            if action == 'ensure_membership' and calls == 3:
                raise portal.PortalError('lost response', uncertain=True)
            return remote(action, data, method)
        api.side_effect = uncertain
        with self.assertRaises(portal.PortalError):
            roster.save_student(self.group.pk, row())
        roster.save_student(self.group.pk, row())
        self.assertEqual(Student.objects.count(), 1)

    @patch('students.portal.call', side_effect=remote)
    def test_local_database_failure_then_retry(self, api):
        with patch('students.roster.Student.objects.update_or_create', side_effect=OperationalError('private database detail')):
            with self.assertRaisesMessage(portal.PortalError, 'Локальное сохранение'):
                roster.save_student(self.group.pk, row())
        roster.save_student(self.group.pk, row())
        self.assertEqual(Student.objects.count(), 1)

    @patch('students.portal.call')
    def test_other_group_rejected_before_network(self, api):
        student = self.student()
        other = Group.objects.create(name='Другая', code='OTHER')
        with self.assertRaises(portal.PortalError):
            roster.save_student(other.pk, row())
        with self.assertRaises(portal.PortalError):
            roster.remove_student(other.pk, student.pk)
        api.assert_not_called()

    @patch('students.portal.call')
    def test_group_with_students_cannot_be_deleted(self, api):
        self.student()
        response = self.client.post(reverse('students:group_delete', args=[self.group.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Group.objects.filter(pk=self.group.pk).exists())
        api.assert_not_called()

    @patch('students.portal.call', side_effect=portal.PortalError('portal unavailable'))
    def test_group_error_visible_and_no_local_save(self, api):
        response = self.client.post(reverse('students:group_create'), {'name': 'New', 'code': 'NEW'})
        self.assertContains(response, 'portal unavailable')
        self.assertFalse(Group.objects.filter(code='NEW').exists())

    @patch('students.portal.call', side_effect=remote)
    def test_group_update_immutable_code(self, api):
        form = GroupForm({'name': 'Новое', 'code': 'CHANGED', 'curators': []}, instance=self.group)
        self.assertTrue(form.is_valid(), form.errors)
        roster.save_group(form)
        self.group.refresh_from_db()
        self.assertEqual(self.group.code, 'RU01')
        self.assertEqual(self.group.name, 'Новое')

    def test_permissions_and_csrf(self):
        target = reverse('students:student_create', args=[self.group.pk])
        self.client.force_login(get_user_model().objects.create_user(username='ordinary'))
        self.assertEqual(self.client.post(target, row()).status_code, 403)
        self.assertEqual(self.client.get(reverse('students:student_import_template', args=[self.group.pk])).status_code, 403)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.staff)
        self.assertEqual(csrf_client.post(target, row()).status_code, 403)

    @patch('students.portal.call', side_effect=remote)
    def test_modal_and_exclusion_post(self, api):
        create = reverse('students:student_create', args=[self.group.pk])
        self.assertContains(self.client.get(create), 'Добавление студента')
        self.assertTrue(self.client.post(create, row()).json()['ok'])
        student = Student.objects.get(code='001')
        edit = reverse('students:student_edit', args=[self.group.pk, student.pk])
        self.assertTrue(self.client.post(edit, row(first_name='Пётр', code='CHANGED')).json()['ok'])
        student.refresh_from_db()
        self.assertEqual(student.code, '001')
        response = self.client.post(reverse('students:remove_student_from_group', args=[self.group.pk, student.pk]))
        self.assertEqual(response.status_code, 302)
        student.refresh_from_db()
        self.assertIsNone(student.group_id)

    def test_all_management_pages_render(self):
        self.student()
        for name in ('group_students', 'student_import', 'student_import_template'):
            self.assertEqual(self.client.get(reverse('students:' + name, args=[self.group.pk])).status_code, 200)

    def test_template_has_text_codes(self):
        book = load_workbook(BytesIO(template_bytes()))
        self.assertEqual([cell.value for cell in book.active[1]], HEADERS)
        self.assertEqual(book.active['E2'].number_format, '@')
        book.close()
        rows = read_rows(self.file().read(), self.group)
        self.assertEqual(rows[0]['code'], '001')

    def test_invalid_workbooks(self):
        invalid = [
            [['Иванов', 'Иван', '', '', '001', '002', '']],
            [['Иванов', 'Иван', '', '', 1, 1, '']],
            [['=1+1', 'Иван', '', '', '001', '001', '']],
            [['Иванов', 'Иван', '', 'bad-email', '001', '001', '']],
            [['Иванов', 'Иван', '', '', '001', '001', '']] * 2,
        ]
        for data in invalid:
            with self.subTest(data=data), self.assertRaises(ImportError):
                read_rows(workbook(data), self.group)
        with self.assertRaises(ImportError):
            read_rows(b'not a spreadsheet', self.group)

    @patch('students.management_views.save_student')
    def test_preview_no_writes_and_tamper_protection(self, save):
        url = reverse('students:student_import', args=[self.group.pk])
        preview = self.client.post(url, {'file': self.file()})
        self.assertEqual(preview.status_code, 200)
        self.assertNotIn('private-password', preview.content.decode())
        save.assert_not_called()
        changed = [['Иванов', 'Иван', '', '', '002', '002', '']]
        response = self.client.post(url, {'phase': 'execute', 'token': preview.json()['token'], 'file': self.file(changed)})
        self.assertEqual(response.status_code, 400)
        save.assert_not_called()

    @patch('students.management_views.save_student')
    def test_import_batches_and_partial_failure(self, save):
        save.side_effect = [(None, 'created'), portal.PortalError('lost', uncertain=True)]
        data = [['Иванов', 'Иван', '', '', f'{i:03}', f'{i:03}', 'secret'] for i in range(26)]
        url = reverse('students:student_import', args=[self.group.pk])
        content = workbook(data)
        token = self.client.post(url, {'file': SimpleUploadedFile('students.xlsx', content)}).json()['token']
        result = self.client.post(url, {'phase': 'execute', 'token': token, 'file': SimpleUploadedFile('students.xlsx', content)}).json()
        self.assertEqual(len(result['results']), 25)
        self.assertEqual(result['results'][0]['status'], 'created')
        self.assertEqual(result['results'][1]['status'], 'uncertain')
        self.assertEqual(result['results'][2]['status'], 'not_run')
        self.assertEqual(result['remaining'], 1)
        self.assertNotIn('secret', str(result))
        self.assertEqual(save.call_count, 2)


@override_settings(PORTAL_ROSTER_URL='https://api.example.test/api/index.html', PORTAL_ROSTER_TOKEN='private-token')
class PortalAdapterTests(TestCase):
    def response(self, body):
        context = Mock()
        context.__enter__ = Mock(return_value=Mock(read=Mock(return_value=body)))
        context.__exit__ = Mock(return_value=False)
        return context

    @patch('students.portal.build_opener')
    def test_string_ids_and_auth_header(self, build):
        build.return_value.open.return_value = self.response(b'{"ok":true,"data":{"group_id":"9000000000000000001","status":"created"}}')
        result = portal.call('upsert_group', {'code': '001'})
        self.assertEqual(result['group_id'], '9000000000000000001')
        request = build.return_value.open.call_args.args[0]
        self.assertEqual(request.get_header('Authorization'), 'private-token')

    @override_settings(PORTAL_ROSTER_URL='https://api.example.test/_wt')
    @patch('students.portal.build_opener')
    def test_webmode_action_route(self, build):
        build.return_value.open.return_value = self.response(b'{"ok":true,"data":{"group_id":"800","status":"created"}}')
        portal.call('upsert_group', {'code': '001'})
        request = build.return_value.open.call_args.args[0]
        self.assertEqual(request.full_url, 'https://api.example.test/_wt/russian_language/action/upsert_group')

    @override_settings(PORTAL_ROSTER_URL='http://127.0.0.1/_wt')
    @patch('students.portal.build_opener')
    def test_loopback_http_for_local_development(self, build):
        build.return_value.open.return_value = self.response(b'{"ok":true,"data":{"group_id":"800","status":"created"}}')
        portal.call('upsert_group', {'code': '001'})
        self.assertEqual(build.return_value.open.call_args.args[0].full_url, 'http://127.0.0.1/_wt/russian_language/action/upsert_group')

    @override_settings(PORTAL_ROSTER_URL='http://api.example.test/_wt')
    @patch('students.portal.build_opener')
    def test_external_plain_http_is_rejected(self, build):
        with self.assertRaises(portal.PortalError):
            portal.call('upsert_group', {})
        build.assert_not_called()

    @patch('students.portal.build_opener')
    def test_invalid_json_and_non_string_id(self, build):
        for body in (b'html error', b'{"ok":true,"data":{"group_id":123,"status":"created"}}', b'{"ok":true,"data":{"status":"created"}}'):
            build.return_value.open.return_value = self.response(body)
            with self.assertRaises(portal.PortalError):
                portal.call('upsert_group', {})

    @patch('students.portal.build_opener')
    def test_errors_do_not_expose_remote_details(self, build):
        for error in (HTTPError('https://example.test', 409, 'private-token', {}, None), URLError('private-token'), TimeoutError('private-token')):
            build.return_value.open.side_effect = error
            with self.assertRaises(portal.PortalError) as caught:
                portal.call('upsert_group', {})
            self.assertNotIn('private-token', str(caught.exception))

    @override_settings(PORTAL_ROSTER_TOKEN='')
    def test_missing_config(self):
        with self.assertRaises(portal.PortalError):
            portal.call('upsert_group', {})


class RosterMigrationTests(TransactionTestCase):
    def test_existing_roster_survives_additive_migration(self):
        executor = MigrationExecutor(connection)
        executor.migrate([('students', '0005_group_curators_m2m')])
        old_apps = executor.loader.project_state([('students', '0005_group_curators_m2m')]).apps
        old_group = old_apps.get_model('students', 'Group').objects.create(name='До миграции', code='OLD-GR')
        old_student = old_apps.get_model('students', 'Student').objects.create(code='0001', first_name='Иван', last_name='Иванов', group_id=old_group.pk)
        try:
            executor = MigrationExecutor(connection)
            executor.migrate([('students', '0006_portal_roster')])
            student = Student.objects.get(pk=old_student.pk)
            self.assertEqual(student.group_id, old_group.pk)
            self.assertEqual(student.code, '0001')
            self.assertIsNone(student.portal_id)
            self.assertEqual(student.email, '')
            self.assertEqual(Group.objects.get(pk=old_group.pk).code, 'OLD-GR')
        finally:
            executor = MigrationExecutor(connection)
            executor.migrate(executor.loader.graph.leaf_nodes())
