"""Validated XLSX input; this module never writes to either system."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .forms import StudentForm
from .models import Student

HEADERS = ['Фамилия', 'Имя', 'Отчество', 'Почта', 'code', 'login', 'password']
KEYS = ['last_name', 'first_name', 'middle_name', 'email', 'code', 'login', 'password']
MAX_SIZE = 5 * 1024 * 1024


class ImportError(ValueError):
    pass


def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Студенты'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    for col in ('A', 'B', 'C', 'D', 'E', 'F', 'G'):
        sheet.column_dimensions[col].width = 24
    for row in sheet.iter_rows(min_row=2, max_row=1001, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = '@'
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def read_rows(content, group):
    if not content or len(content) > MAX_SIZE:
        raise ImportError('Выберите XLSX размером не более 5 МБ.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ImportError('Распакованный файл слишком велик.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, OSError, ValueError, KeyError, InvalidFileException, ParseError):
        raise ImportError('Не удалось прочитать XLSX.') from None
    rows, errors, seen = [], [], set()
    try:
        sheet = book.active
        if sheet.max_row and sheet.max_row > 1001:
            raise ImportError('Допускается не более 1000 строк студентов.')
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1, max_col=20))
        headers = [str(cell.value or '').strip() for cell in header_cells]
        if any(headers.count(header) != 1 for header in HEADERS):
            raise ImportError('Нужны уникальные заголовки: ' + ', '.join(HEADERS))
        indices = [headers.index(header) for header in HEADERS]
        existing = dict(Student.objects.values_list('code', 'group_id'))
        for number, cells in enumerate(sheet.iter_rows(min_row=2, max_row=1001, max_col=20), start=2):
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == 'f' for cell in cells):
                errors.append(f'Строка {number}: формулы не допускаются.')
                continue
            data = {key: str(cells[index].value or '').strip() for key, index in zip(KEYS, indices)}
            # Excel numeric codes may have already lost leading zeros; require text.
            if any(not isinstance(cells[indices[index]].value, str) for index in (4, 5)):
                errors.append(f'Строка {number}: code и login должны быть текстом.')
            form = StudentForm(data)
            if not form.is_valid():
                for field, messages in form.errors.items():
                    errors.append(f'Строка {number}, {form.fields[field].label}: ' + ' '.join(messages))
                continue
            if data['code'] != data['login']:
                errors.append(f'Строка {number}: code должен совпадать с login.')
            if data['code'] in seen:
                errors.append(f'Строка {number}: повтор кода {data["code"]}.')
            seen.add(data['code'])
            if existing.get(data['code']) not in (None, group.pk):
                errors.append(f'Строка {number}: студент уже состоит в другой группе.')
            rows.append({'row_number': number, **form.cleaned_data, 'expected': 'обновление' if data['code'] in existing else 'создание / поиск на портале'})
    except ImportError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, ParseError, TypeError, IndexError, StopIteration):
        raise ImportError('Ошибка чтения строк XLSX. Сохраните файл заново по шаблону.') from None
    finally:
        book.close()
    if errors:
        raise ImportError('\n'.join(errors[:50]))
    if not rows:
        raise ImportError('В файле нет студентов.')
    return rows
